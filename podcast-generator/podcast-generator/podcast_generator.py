"""Document-to-podcast generator.

The generator turns a PDF, TXT, or DOCX document into a two-speaker podcast
script with Sarah and Mike, then optionally renders it to MP3.
"""

from __future__ import annotations

import os
import re
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable

import PyPDF2
from dotenv import load_dotenv
from openai import OpenAI

import config
from heygen_client import HeyGenClient, download_audio

load_dotenv()

ProgressCallback = Callable[[str], None]


@dataclass(frozen=True)
class DialogueSegment:
    """One speaker turn in the generated podcast script."""

    speaker: str
    text: str


class PodcastGenerator:
    """Generate Mike/Sarah podcast scripts and audio from local documents."""

    def __init__(
        self,
        openai_api_key: str | None = None,
        script_model: str | None = None,
        tts_model: str | None = None,
        speaker_names: dict[str, str] | None = None,
        speaker_voices: dict[str, str] | None = None,
        speaker_voice_providers: dict[str, str] | None = None,
        heygen_api_key: str | None = None,
        progress_callback: ProgressCallback | None = None,
    ):
        self.api_key = openai_api_key or os.getenv("OPENAI_API_KEY")
        if not self.api_key:
            raise ValueError(
                "OpenAI API key not provided. Set OPENAI_API_KEY or pass openai_api_key."
            )

        self.client = OpenAI(api_key=self.api_key)
        self.script_model = script_model or os.getenv("OPENAI_SCRIPT_MODEL", config.SCRIPT_MODEL)
        self.tts_model = tts_model or os.getenv("OPENAI_TTS_MODEL", config.TTS_MODEL)
        self.speaker_names = self._normalize_speaker_names(speaker_names or config.SPEAKER_NAMES)
        self.voices = dict(config.SPEAKER_VOICES)
        if speaker_voices:
            self.voices.update(speaker_voices)
        self.voice_providers = self._normalize_voice_providers(speaker_voice_providers)
        self.heygen_api_key = heygen_api_key
        self.words_per_minute = config.WORDS_PER_MINUTE
        self.target_duration_minutes = config.DEFAULT_DURATION_MINUTES
        self.progress_callback = progress_callback

    def _log(self, message: str) -> None:
        if self.progress_callback:
            self.progress_callback(message)
        print(message)

    @staticmethod
    def _normalize_speaker_names(names: dict[str, str]) -> dict[str, str]:
        normalized = dict(config.SPEAKER_NAMES)
        for speaker in ("woman", "man"):
            value = (names.get(speaker) or normalized[speaker]).strip()
            value = re.sub(r"[^A-Za-z0-9 '\-]", "", value).strip()
            normalized[speaker] = (value or config.SPEAKER_NAMES[speaker]).upper()
        return normalized

    def _speaker_label(self, speaker: str) -> str:
        return self.speaker_names.get(speaker, config.SPEAKER_NAMES.get(speaker, speaker.upper()))

    def _speaker_label_pattern(self) -> str:
        labels = {
            self._speaker_label("woman"),
            self._speaker_label("man"),
            "SARAH",
            "MIKE",
        }
        escaped = sorted((re.escape(label) for label in labels), key=len, reverse=True)
        return "|".join(escaped)

    def _speaker_from_label(self, label: str) -> str | None:
        normalized = label.strip().upper()
        if normalized in {self._speaker_label("woman"), "SARAH"}:
            return "woman"
        if normalized in {self._speaker_label("man"), "MIKE"}:
            return "man"
        return None

    @staticmethod
    def _normalize_voice_providers(providers: dict[str, str] | None) -> dict[str, str]:
        normalized = {"woman": "openai", "man": "openai"}
        for speaker, provider in (providers or {}).items():
            provider_value = provider.strip().lower()
            if provider_value not in {"openai", "heygen"}:
                raise ValueError("Voice provider must be 'openai' or 'heygen'.")
            normalized[speaker] = provider_value
        return normalized

    @staticmethod
    def validate_duration(duration_minutes: int) -> int:
        """Validate and normalize the requested duration preset."""
        try:
            duration = int(duration_minutes)
        except (TypeError, ValueError) as exc:
            raise ValueError("Duration must be one of: 5, 10, 15, 20 minutes.") from exc

        if duration not in config.VALID_DURATIONS_MINUTES:
            allowed = ", ".join(str(value) for value in config.VALID_DURATIONS_MINUTES)
            raise ValueError(f"Duration must be one of: {allowed} minutes.")
        return duration

    def target_word_count(self, duration_minutes: int) -> int:
        return self.validate_duration(duration_minutes) * self.words_per_minute

    def word_count_bounds(self, duration_minutes: int) -> tuple[int, int]:
        target = self.target_word_count(duration_minutes)
        tolerance = config.SCRIPT_WORD_TOLERANCE
        return round(target * (1 - tolerance)), round(target * (1 + tolerance))

    @staticmethod
    def count_words(text: str) -> int:
        return len(re.findall(r"\b[\w'-]+\b", text))

    def extract_text(self, document_path: str | Path) -> str:
        """Extract text from a supported document type."""
        path = Path(document_path)
        if not path.exists():
            raise FileNotFoundError(f"Document not found: {path}")

        suffix = path.suffix.lower()
        self._log(f"Extracting text from {path.name}...")

        if suffix == ".pdf":
            text = self.extract_text_from_pdf(path)
        elif suffix == ".txt":
            text = self.extract_text_from_txt(path)
        elif suffix == ".docx":
            text = self.extract_text_from_docx(path)
        else:
            raise ValueError("Unsupported document type. Upload a PDF, TXT, or DOCX file.")

        normalized = self._normalize_text(text)
        if not normalized:
            raise ValueError("No readable text was found in the document.")

        self._log(f"Extracted {len(normalized):,} characters.")
        return normalized

    def extract_text_from_pdf(self, pdf_path: str | Path) -> str:
        """Extract text from a PDF file, preferring pdfplumber when available."""
        path = Path(pdf_path)
        try:
            import pdfplumber

            parts: list[str] = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        parts.append(page_text)
            return "\n".join(parts)
        except ImportError:
            return self._extract_text_from_pdf_pypdf2(path)
        except Exception as exc:
            raise RuntimeError(f"Error extracting PDF text from {path.name}: {exc}") from exc

    @staticmethod
    def _extract_text_from_pdf_pypdf2(path: Path) -> str:
        parts: list[str] = []
        with path.open("rb") as file:
            reader = PyPDF2.PdfReader(file)
            for page in reader.pages:
                page_text = page.extract_text() or ""
                if page_text.strip():
                    parts.append(page_text)
        return "\n".join(parts)

    @staticmethod
    def extract_text_from_txt(txt_path: str | Path) -> str:
        path = Path(txt_path)
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                return path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        raise RuntimeError(f"Could not decode text file: {path.name}")

    @staticmethod
    def extract_text_from_docx(docx_path: str | Path) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX support requires python-docx. Install requirements.txt.") from exc

        document = Document(str(docx_path))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)
        return "\n".join(parts)

    @staticmethod
    def _normalize_text(text: str) -> str:
        text = text.replace("\x00", " ")
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    @staticmethod
    def chunk_text(text: str, max_chars: int = config.SOURCE_CHUNK_CHARS) -> list[str]:
        """Chunk text on paragraph boundaries where possible."""
        paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
        chunks: list[str] = []
        current = ""

        for paragraph in paragraphs:
            if len(paragraph) > max_chars:
                if current:
                    chunks.append(current.strip())
                    current = ""
                chunks.extend(PodcastGenerator._hard_wrap(paragraph, max_chars))
                continue

            candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
            if len(candidate) <= max_chars:
                current = candidate
            else:
                chunks.append(current.strip())
                current = paragraph

        if current:
            chunks.append(current.strip())
        return chunks

    @staticmethod
    def _hard_wrap(text: str, max_chars: int) -> list[str]:
        return [text[index : index + max_chars].strip() for index in range(0, len(text), max_chars)]

    def synthesize_source_notes(self, document_text: str, paper_title: str) -> str:
        """Create source-grounded notes that later script prompts can reuse."""
        self._log("Synthesizing source notes...")
        clipped_text = document_text[: config.MAX_SOURCE_CHARS_FOR_NOTES]
        chunks = self.chunk_text(clipped_text, config.SOURCE_CHUNK_CHARS)
        notes: list[str] = []

        for index, chunk in enumerate(chunks, start=1):
            self._log(f"Summarizing source chunk {index}/{len(chunks)}...")
            prompt = (
                "Extract source-grounded notes for a podcast script. Focus on the document's "
                "main topic, methods or argument, evidence, findings, conclusions, named "
                "concepts, and any implications for future risks, risk management, roles, "
                "assurance, and governance. Do not invent facts. Use concise bullets.\n\n"
                f"Title: {paper_title}\n\nDocument excerpt:\n{chunk}"
            )
            notes.append(
                self._call_text_model(
                    system="You extract accurate notes from source documents.",
                    user=prompt,
                    max_output_tokens=1300,
                )
            )

        combined = "\n\n".join(notes)
        return self._normalize_text(combined)

    def generate_podcast_script(
        self,
        document_text: str,
        paper_title: str = "Document",
        duration_minutes: int = config.DEFAULT_DURATION_MINUTES,
    ) -> list[dict[str, str]]:
        """Generate and validate a Mike/Sarah dialogue script."""
        duration = self.validate_duration(duration_minutes)
        self.target_duration_minutes = duration
        target_words = self.target_word_count(duration)
        source_notes = self.synthesize_source_notes(document_text, paper_title)

        self._log(f"Generating {duration}-minute script target ({target_words} words)...")
        section_texts: list[str] = []
        previous_summary = ""

        for section_name, section_goal, fraction in config.SCRIPT_SECTIONS:
            section_words = max(60, round(target_words * fraction))
            self._log(f"Writing section: {section_name} (~{section_words} words)...")
            section_text = self._generate_script_section(
                paper_title=paper_title,
                source_notes=source_notes,
                section_name=section_name,
                section_goal=section_goal,
                target_words=section_words,
                previous_summary=previous_summary,
                duration_minutes=duration,
            )
            section_texts.append(section_text)
            previous_summary = self._summarize_previous_script(section_texts)

        script_text = self._normalize_script_text("\n\n".join(section_texts))
        dialogue = self._parse_script(script_text)
        dialogue = self._repair_dialogue_if_needed(dialogue, source_notes, paper_title, duration)

        words = self._dialogue_word_count(dialogue)
        estimated = words / self.words_per_minute
        self._log(f"Generated {words} words; estimated duration {estimated:.1f} minutes.")
        return [segment.__dict__ for segment in dialogue]

    def _generate_script_section(
        self,
        paper_title: str,
        source_notes: str,
        section_name: str,
        section_goal: str,
        target_words: int,
        previous_summary: str,
        duration_minutes: int,
    ) -> str:
        woman_label = self._speaker_label("woman")
        man_label = self._speaker_label("man")
        system = (
            "You are a senior podcast script writer. Write only dialogue lines. "
            f"Every line must start with {woman_label}: or {man_label}:. "
            f"{woman_label.title()} is an audit and risk management expert. "
            f"{man_label.title()} is a governance specialist and researcher. "
            "The dialogue must sound like a real, curious podcast conversation: "
            "responsive, source-grounded, intellectually rigorous, and accessible. "
            "Avoid lecture-like monologues."
        )
        user = f"""
Create the {section_name} section of a two-speaker podcast about "{paper_title}".

Section goal: {section_goal}
Overall episode target: {duration_minutes} minutes.
This section target: about {target_words} spoken words.

Rules:
- Write only speaker dialogue, one turn per paragraph.
- Alternate naturally between {woman_label} and {man_label}, but make each turn respond to the previous one.
- Use the source notes only; do not invent study details, numbers, authors, or claims.
- Explain acronyms in full and avoid acronym-heavy speech.
- Keep {woman_label.title()} and {man_label.title()} balanced across the section.
- Do not include markdown headings, timestamps, music cues, or narration.
- Make the section self-contained but continuous with the previous section.
- Keep most turns to 1-3 spoken sentences. Avoid long mini-lectures.
- Include natural interaction: follow-up questions, brief reactions, clarifications, and callbacks to earlier points.
- Let them occasionally challenge or refine each other's interpretation politely.
- Avoid repeatedly saying each other's names. Use names only when it sounds natural.
- Do not simply alternate summaries. Mike and Sarah should build on each other's ideas.

Previous section summary:
{previous_summary or "This is the first section."}

Source notes:
{source_notes}
"""
        return self._call_text_model(
            system=system,
            user=user,
            max_output_tokens=max(300, round(target_words * 1.7)),
        )

    def _summarize_previous_script(self, section_texts: Iterable[str]) -> str:
        text = "\n\n".join(section_texts)
        words = text.split()
        if len(words) <= 220:
            return text
        return " ".join(words[-220:])

    def _repair_dialogue_if_needed(
        self,
        dialogue: list[DialogueSegment],
        source_notes: str,
        paper_title: str,
        duration_minutes: int,
    ) -> list[DialogueSegment]:
        if not dialogue:
            raise RuntimeError("The model did not return any Mike/Sarah dialogue.")

        min_words, max_words = self.word_count_bounds(duration_minutes)
        words = self._dialogue_word_count(dialogue)
        if min_words <= words <= max_words:
            return dialogue

        target_words = self.target_word_count(duration_minutes)
        current_dialogue = dialogue
        best_dialogue = dialogue
        best_words = words

        woman_label = self._speaker_label("woman")
        man_label = self._speaker_label("man")

        for attempt in range(1, 4):
            current_words = self._dialogue_word_count(current_dialogue)
            if abs(current_words - target_words) < abs(best_words - target_words):
                best_dialogue = current_dialogue
                best_words = current_words
            if min_words <= current_words <= max_words:
                return current_dialogue

            action = "expand" if current_words < min_words else "condense"
            self._log(
                f"Script is {current_words} words; repair pass {attempt} to "
                f"{target_words} words ({min_words}-{max_words} acceptable)."
            )

            current_script = self.format_script(current_dialogue)
            system = (
                "You revise podcast scripts. Return only dialogue lines that start with "
                f"{woman_label}: or {man_label}:. Keep the script source-grounded and make it sound "
                "like a natural podcast conversation, not alternating mini-lectures."
            )
            user = f"""
Revise this complete Mike/Sarah podcast script about "{paper_title}".

Current word count: {current_words}.
Required final range: {min_words} to {max_words} spoken words.
Ideal target: about {target_words} spoken words.
Task: {action} the script so the final answer is inside the required range.

Hard requirements:
- Return the complete revised script, not just edits.
- Do not exceed {max_words} words.
- Do not go below {min_words} words.
- Preserve the flow: understand the document, analyze implications, close with takeaways.
- Keep {woman_label.title()} and {man_label.title()} balanced.
- Use only the source notes for factual claims.
- Avoid acronym-heavy wording and explain necessary terms in full.
- Return only {woman_label}: and {man_label}: dialogue lines.
- Keep most turns short: 1-3 spoken sentences.
- For short episodes, use compact exchanges. Most turns should stay under 45 words.
- Add more direct interaction: follow-up questions, reactions, clarifications, callbacks, and polite refinements.
- Avoid repeated name callouts and stiff transitions.
- Preserve factual accuracy while making the exchange more conversational.
- If the script is too long, remove repeated framing, extra examples, and redundant transitions before removing source facts.

Source notes:
{source_notes}

Current script:
{current_script}
"""
            if action == "condense":
                token_budget = max(500, round(max_words * 1.45))
            else:
                token_budget = max(900, round(target_words * 2.2))
            repaired_text = self._call_text_model(
                system=system,
                user=user,
                max_output_tokens=token_budget,
            )
            repaired_dialogue = self._parse_script(repaired_text)
            if not repaired_dialogue:
                raise RuntimeError("Script repair failed: no valid Mike/Sarah dialogue returned.")
            current_dialogue = repaired_dialogue
            repaired_words = self._dialogue_word_count(current_dialogue)
            if abs(repaired_words - target_words) < abs(best_words - target_words):
                best_dialogue = current_dialogue
                best_words = repaired_words

        if best_words > max_words:
            trimmed_dialogue = self._trim_dialogue_to_word_limit(best_dialogue, max_words)
            trimmed_words = self._dialogue_word_count(trimmed_dialogue)
            if min_words <= trimmed_words <= max_words:
                self._log(f"Trimmed script to {trimmed_words} words after repair.")
                return trimmed_dialogue

        if best_words < min_words:
            expanded_dialogue = self._expand_dialogue_to_word_minimum(
                best_dialogue,
                min_words=min_words,
                max_words=max_words,
            )
            expanded_words = self._dialogue_word_count(expanded_dialogue)
            if min_words <= expanded_words <= max_words:
                self._log(f"Expanded script to {expanded_words} words after repair.")
                return expanded_dialogue

        final_words = self._dialogue_word_count(best_dialogue)
        self._log(f"Warning: repaired script is {final_words} words, outside the target range.")
        return best_dialogue

    def _trim_dialogue_to_word_limit(
        self,
        dialogue: list[DialogueSegment],
        max_words: int,
    ) -> list[DialogueSegment]:
        """Trim longest turns just enough to satisfy a word limit."""
        trimmed = [DialogueSegment(segment.speaker, segment.text) for segment in dialogue]

        while self._dialogue_word_count(trimmed) > max_words:
            longest_index = max(
                range(len(trimmed)),
                key=lambda index: self.count_words(trimmed[index].text),
            )
            segment = trimmed[longest_index]
            words = segment.text.split()
            if len(words) <= 12:
                break

            remove_count = min(self._dialogue_word_count(trimmed) - max_words, max(1, len(words) // 4))
            new_words = words[:-remove_count]
            new_text = " ".join(new_words).strip()
            if new_text and new_text[-1] not in ".!?":
                new_text += "."
            trimmed[longest_index] = DialogueSegment(segment.speaker, new_text)

        return trimmed

    def _expand_dialogue_to_word_minimum(
        self,
        dialogue: list[DialogueSegment],
        min_words: int,
        max_words: int,
    ) -> list[DialogueSegment]:
        """Add compact conversational turns until the script reaches a word minimum."""
        expanded = [DialogueSegment(segment.speaker, segment.text) for segment in dialogue]
        additions = [
            (
                "woman",
                "Before we move on, one practical point is worth underlining: the value here is not only in the tool itself, but in how clearly people define responsibility around it.",
            ),
            (
                "man",
                "That is a useful distinction. The document keeps coming back to accountability, because without named owners, even strong technical systems can become hard to review or correct.",
            ),
            (
                "woman",
                "So the listener should hear this as a governance lesson as much as a technology lesson: evidence, privacy, review, and escalation all need to be visible.",
            ),
            (
                "man",
                "And that visibility matters for assurance. If a team cannot explain what happened, who approved it, and what risks were tracked, trust becomes much harder to maintain.",
            ),
            (
                "woman",
                "Exactly. A good framework makes those questions routine rather than exceptional, which is what helps research teams use new systems responsibly.",
            ),
            (
                "man",
                "In other words, the future opportunity is real, but the controls have to grow with it. That balance is the thread running through the whole discussion.",
            ),
        ]

        addition_index = 0
        while self._dialogue_word_count(expanded) < min_words and addition_index < len(additions):
            speaker, text = additions[addition_index]
            expanded.append(DialogueSegment(speaker, text))
            addition_index += 1

        words = self._dialogue_word_count(expanded)
        if words > max_words:
            expanded = self._trim_dialogue_to_word_limit(expanded, max_words)

        return expanded

    def _call_text_model(self, system: str, user: str, max_output_tokens: int) -> str:
        budgets = [max_output_tokens]
        larger_budget = max(max_output_tokens * 2, 4000)
        if larger_budget != max_output_tokens:
            budgets.append(larger_budget)

        last_status = "unknown"
        for budget in budgets:
            response = self.client.responses.create(
                model=self.script_model,
                input=[
                    {"role": "system", "content": system},
                    {"role": "user", "content": user},
                ],
                max_output_tokens=budget,
            )
            last_status = getattr(response, "status", "unknown")
            text = self._extract_response_text(response)
            if text:
                return text

        raise RuntimeError(f"OpenAI returned an empty text response; response status was {last_status}.")

    @staticmethod
    def _extract_response_text(response) -> str:
        output_text = getattr(response, "output_text", None)
        if output_text:
            return output_text.strip()

        pieces: list[str] = []
        for item in getattr(response, "output", []) or []:
            for content in getattr(item, "content", []) or []:
                text = getattr(content, "text", None)
                if text:
                    pieces.append(text)
        return "\n".join(pieces).strip()

    @staticmethod
    def _normalize_script_text_static(script_text: str, label_pattern: str) -> str:
        lines = []
        for raw_line in script_text.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            if re.match(rf"^({label_pattern})\s*:", line, flags=re.IGNORECASE):
                lines.append(line)
        return "\n".join(lines)

    def _normalize_script_text(self, script_text: str) -> str:
        return self._normalize_script_text_static(script_text, self._speaker_label_pattern())

    def _parse_script(self, script_text: str) -> list[DialogueSegment]:
        """Parse raw model output into speaker turns."""
        dialogue: list[DialogueSegment] = []
        current_speaker: str | None = None
        current_text: list[str] = []

        def flush() -> None:
            nonlocal current_speaker, current_text
            text = " ".join(part.strip() for part in current_text if part.strip()).strip()
            if current_speaker and text:
                dialogue.append(DialogueSegment(speaker=current_speaker, text=text))
            current_speaker = None
            current_text = []

        for raw_line in script_text.splitlines():
            line = raw_line.strip()
            if not line:
                continue

            match = re.match(
                rf"^({self._speaker_label_pattern()})\s*:\s*(.+)$",
                line,
                flags=re.IGNORECASE,
            )
            if match:
                flush()
                current_speaker = self._speaker_from_label(match.group(1))
                current_text = [match.group(2).strip()]
            elif re.match(r"^[A-Z][A-Z0-9 _-]{1,40}\s*:", line):
                flush()
            elif current_speaker:
                current_text.append(line)

        flush()
        return dialogue

    def format_script(self, dialogue: Iterable[DialogueSegment | dict[str, str]]) -> str:
        lines: list[str] = []
        for segment in dialogue:
            if isinstance(segment, dict):
                speaker = segment["speaker"]
                text = segment["text"]
            else:
                speaker = segment.speaker
                text = segment.text
            speaker_name = self._speaker_label(speaker)
            lines.append(f"{speaker_name}: {text}")
        return "\n\n".join(lines)

    def _dialogue_word_count(self, dialogue: Iterable[DialogueSegment]) -> int:
        return sum(self.count_words(segment.text) for segment in dialogue)

    def save_script(self, dialogue: list[dict[str, str]], output_path: str | Path) -> None:
        """Save the podcast script to a text file."""
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        segments = [DialogueSegment(**segment) for segment in dialogue]
        content = "PODCAST SCRIPT\n" + "=" * 50 + "\n\n" + self.format_script(segments) + "\n"
        path.write_text(content, encoding="utf-8")
        self._log(f"Script saved: {path}")

    def generate_audio(self, dialogue: list[dict[str, str]], output_path: str | Path) -> None:
        """Generate MP3 audio from dialogue using chunked TTS requests."""
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self._log("Generating audio...")

        segments = [DialogueSegment(**segment) for segment in dialogue]
        audio_paths: list[Path] = []

        with tempfile.TemporaryDirectory(prefix="podcast_segments_") as temp_dir:
            temp_root = Path(temp_dir)
            segment_counter = 0
            for index, segment in enumerate(segments, start=1):
                parts = self.split_tts_input(segment.text)
                for part_index, part in enumerate(parts, start=1):
                    segment_counter += 1
                    self._log(
                        f"TTS segment {segment_counter} "
                        f"({index}/{len(segments)}, part {part_index}/{len(parts)})..."
                    )
                    audio_path = temp_root / f"segment_{segment_counter:04d}.mp3"
                    self._write_tts_segment(segment.speaker, part, audio_path)
                    if audio_path.stat().st_size <= 0:
                        raise RuntimeError(f"TTS returned an empty audio segment: {audio_path.name}")
                    audio_paths.append(audio_path)

            if not audio_paths:
                raise RuntimeError("No audio segments were generated.")

            self._combine_audio_segments(audio_paths, path)

        self._log(f"Audio saved: {path}")

    def _write_tts_segment(self, speaker: str, text: str, output_path: Path) -> None:
        provider = self.voice_providers.get(speaker, "openai")
        if provider == "heygen":
            self._write_heygen_tts_segment(speaker, text, output_path)
            return

        kwargs = {
            "model": self.tts_model,
            "voice": self.voices[speaker],
            "input": text,
            "speed": config.SPEECH_SPEED,
            "response_format": "mp3",
        }
        if not self.tts_model.startswith("tts-1"):
            kwargs["instructions"] = self._tts_instructions(speaker)

        response = self.client.audio.speech.create(**kwargs)
        output_path.write_bytes(response.content)

    def _write_heygen_tts_segment(self, speaker: str, text: str, output_path: Path) -> None:
        if not self.heygen_api_key:
            raise ValueError("HeyGen API key is required for HeyGen voices.")
        client = HeyGenClient(self.heygen_api_key)
        audio_url = client.create_speech_url(
            voice_id=self.voices[speaker],
            text=text,
            speed=config.SPEECH_SPEED,
        )
        output_path.write_bytes(download_audio(audio_url))

    def _tts_instructions(self, speaker: str) -> str:
        name = self._speaker_label(speaker).title()
        if speaker == "woman":
            return (
                f"Speak as {name}: warm, thoughtful, clear, and engaged. "
                "Use a professional podcast-host tone."
            )
        return (
            f"Speak as {name}: calm, precise, curious, and conversational. "
            "Use a professional researcher tone."
        )

    @staticmethod
    def split_tts_input(text: str, max_chars: int = config.TTS_MAX_INPUT_CHARS) -> list[str]:
        """Split text into speech endpoint sized chunks."""
        normalized = re.sub(r"\s+", " ", text).strip()
        if len(normalized) <= max_chars:
            return [normalized] if normalized else []

        sentences = re.split(r"(?<=[.!?])\s+", normalized)
        chunks: list[str] = []
        current = ""

        for sentence in sentences:
            if len(sentence) > max_chars:
                if current:
                    chunks.append(current.strip())
                    current = ""
                chunks.extend(PodcastGenerator._hard_wrap(sentence, max_chars))
                continue

            candidate = f"{current} {sentence}".strip() if current else sentence
            if len(candidate) <= max_chars:
                current = candidate
            else:
                chunks.append(current.strip())
                current = sentence

        if current:
            chunks.append(current.strip())
        return chunks

    def _combine_audio_segments(self, segment_paths: list[Path], output_path: Path) -> None:
        try:
            from pydub import AudioSegment
        except ImportError as exc:
            raise RuntimeError("Audio combining requires pydub. Install requirements.txt.") from exc

        combined = AudioSegment.empty()
        pause = AudioSegment.silent(duration=config.PAUSE_BETWEEN_SPEAKERS_MS)
        for segment_path in segment_paths:
            audio = AudioSegment.from_mp3(segment_path)
            combined += audio
            combined += pause
        combined.export(output_path, format="mp3")

        if output_path.stat().st_size <= 0:
            raise RuntimeError(f"Combined audio file is empty: {output_path}")

    def create_podcast(
        self,
        document_path: str | Path | None = None,
        output_dir: str | Path = config.DEFAULT_OUTPUT_DIR,
        paper_title: str | None = None,
        duration_minutes: int = config.DEFAULT_DURATION_MINUTES,
        generate_audio: bool = True,
        pdf_path: str | Path | None = None,
    ) -> tuple[str, str | None]:
        """Create a podcast script and optional audio from a document.

        The ``pdf_path`` argument is retained for backward compatibility; new
        callers should use ``document_path``.
        """
        if document_path is None:
            document_path = pdf_path
        if document_path is None:
            raise ValueError("document_path is required.")

        duration = self.validate_duration(duration_minutes)
        source_path = Path(document_path)
        title = paper_title or source_path.stem
        output_root = Path(output_dir)
        output_root.mkdir(parents=True, exist_ok=True)

        self._log("=" * 60)
        self._log("Podcast generation started")
        self._log(f"Document: {source_path}")
        self._log(f"Title: {title}")
        self._log(f"Duration target: {duration} minutes")
        self._log("=" * 60)

        document_text = self.extract_text(source_path)
        dialogue = self.generate_podcast_script(document_text, title, duration)

        timestamp = datetime.now().strftime(config.TIMESTAMP_FORMAT)
        script_path = output_root / f"script_{timestamp}.txt"
        self.save_script(dialogue, script_path)

        audio_path: Path | None = None
        if generate_audio:
            min_words, max_words = self.word_count_bounds(duration)
            script_words = self._dialogue_word_count(DialogueSegment(**segment) for segment in dialogue)
            if not min_words <= script_words <= max_words:
                raise RuntimeError(
                    f"Audio generation skipped because the script is {script_words} words; "
                    f"required range is {min_words}-{max_words} words. Script saved at {script_path}."
                )
            audio_path = output_root / f"podcast_{timestamp}.mp3"
            self.generate_audio(dialogue, audio_path)

        self._log("=" * 60)
        self._log("Podcast generation complete")
        self._log(f"Script: {script_path}")
        if audio_path:
            self._log(f"Audio: {audio_path}")
        self._log("=" * 60)

        return str(script_path), str(audio_path) if audio_path else None


def main() -> None:
    """Small CLI fallback when running this module directly."""
    import argparse

    parser = argparse.ArgumentParser(description="Generate a Mike/Sarah podcast from a document.")
    parser.add_argument("--input", required=True, help="Path to a PDF, TXT, or DOCX document.")
    parser.add_argument("--output-dir", default=config.DEFAULT_OUTPUT_DIR, help="Output directory.")
    parser.add_argument("--title", default=None, help="Optional podcast/document title.")
    parser.add_argument(
        "--duration",
        type=int,
        default=config.DEFAULT_DURATION_MINUTES,
        choices=config.VALID_DURATIONS_MINUTES,
        help="Podcast duration preset in minutes.",
    )
    parser.add_argument("--script-only", action="store_true", help="Generate script without audio.")
    args = parser.parse_args()

    generator = PodcastGenerator()
    script_path, audio_path = generator.create_podcast(
        document_path=args.input,
        output_dir=args.output_dir,
        paper_title=args.title,
        duration_minutes=args.duration,
        generate_audio=not args.script_only,
    )
    print(f"Script ready: {script_path}")
    if audio_path:
        print(f"Audio ready: {audio_path}")


if __name__ == "__main__":
    main()
